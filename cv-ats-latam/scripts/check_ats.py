#!/usr/bin/env python3
"""
Verifica que un CV (PDF o DOCX) sea legible por un ATS y, opcionalmente, qué
palabras clave de una oferta cubre.

Uso:
    python check_ats.py CV.pdf [--json datos.json] [--keywords "SQL, Power BI, presupuesto"]
                               [--keywords-file claves.txt]

Qué revisa (todo es verificable, no hay "puntajes" inventados):
  1. Que el texto se pueda extraer (no es una imagen escaneada).
  2. Que correo y teléfono aparezcan y estén al inicio.
  3. Que estén las secciones estándar con nombres que un ATS reconoce.
  4. Con --json: que cada cargo, empresa y logro aparezca en el texto extraído
     y en el mismo orden (es decir, que el ATS lo lee como se diseñó).
  5. En PDF: fuentes incrustadas, sin imágenes, número de páginas, sin
     caracteres raros (ligaduras, símbolos que se pierden al copiar).
  6. En DOCX: sin tablas, cuadros de texto, imágenes ni texto en encabezado/pie.
  7. Con --keywords: qué palabras clave de la oferta aparecen y cuáles faltan.
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
import unicodedata
import zipfile

OK, AVISO, FALLA = "✓", "⚠", "✗"


def normalizar(t):
    t = unicodedata.normalize("NFKD", t or "")
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = t.lower().replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", t).strip()


# ----------------------------------------------------------------------------
# Extracción
# ----------------------------------------------------------------------------
def textos_pdf(ruta):
    """Extrae el texto con varios lectores distintos, porque cada ATS usa el suyo
    y no todos reconstruyen el orden de lectura igual. Devuelve {nombre: texto}."""
    textos = {}
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", "-enc", "UTF-8", ruta, "-"], capture_output=True, text=True)
        if r.stdout.strip():
            textos["pdftotext"] = r.stdout
    try:
        from pypdf import PdfReader
        t = "\n".join(p.extract_text() or "" for p in PdfReader(ruta).pages)
        if t.strip():
            textos["pypdf"] = t
    except Exception:  # noqa: BLE001
        pass
    try:
        from pdfminer.high_level import extract_text
        t = extract_text(ruta)
        if t.strip():
            textos["pdfminer"] = t
    except Exception:  # noqa: BLE001
        pass
    if not textos:
        sys.exit("No pude extraer texto: instala poppler-utils (pdftotext), pypdf o pdfminer.six.")
    return textos


def revisar_pdf_tecnico(ruta, res):
    if shutil.which("pdfinfo"):
        info = subprocess.run(["pdfinfo", ruta], capture_output=True, text=True).stdout
        m = re.search(r"Pages:\s+(\d+)", info)
        if m:
            n = int(m.group(1))
            if n <= 2:
                res.append((OK, f"Largo: {n} página(s)."))
            else:
                res.append((AVISO, f"Largo: {n} páginas. En Chile/LatAm lo esperado es 1-2."))
    if shutil.which("pdffonts"):
        out = subprocess.run(["pdffonts", ruta], capture_output=True, text=True).stdout.splitlines()[2:]
        no_emb = [l.split()[0] for l in out if l.strip() and len(l.split()) >= 5 and l.split()[-5] == "no"]
        if no_emb:
            res.append((AVISO, f"Fuentes no incrustadas: {', '.join(no_emb)}. Puede verse distinto en otro equipo."))
        elif out:
            res.append((OK, "Todas las fuentes están incrustadas."))
    if shutil.which("pdfimages"):
        out = subprocess.run(["pdfimages", "-list", ruta], capture_output=True, text=True).stdout.splitlines()[2:]
        n = len([l for l in out if l.strip()])
        if n:
            res.append((AVISO, f"El PDF contiene {n} imagen(es). Un ATS no lee texto dentro de imágenes."))
        else:
            res.append((OK, "Sin imágenes (todo el contenido es texto)."))


def texto_docx(ruta, res):
    try:
        import docx
    except ImportError:
        sys.exit("Necesito python-docx para leer el DOCX: pip install python-docx")
    d = docx.Document(ruta)
    lineas = [p.text for p in d.paragraphs]
    if d.tables:
        res.append((FALLA, f"El documento usa {len(d.tables)} tabla(s). Muchos ATS leen las tablas celda por "
                           "celda y desordenan fechas, cargos y empresas."))
        for t in d.tables:
            for fila in t.rows:
                lineas.append(" ".join(c.text for c in fila.cells))
    else:
        res.append((OK, "Sin tablas."))
    with zipfile.ZipFile(ruta) as z:
        xml = z.read("word/document.xml").decode("utf8", "ignore")
        if "<w:txbxContent" in xml:
            res.append((FALLA, "Hay cuadros de texto. Su contenido suele perderse en el ATS."))
        if "<w:drawing" in xml or "<w:pict" in xml:
            res.append((AVISO, "Hay imágenes o formas (foto, íconos, gráficos). El ATS no las lee."))
        cabeceras = [n for n in z.namelist() if re.match(r"word/(header|footer)\d*\.xml", n)]
        con_texto = []
        for n in cabeceras:
            contenido = re.sub(r"<[^>]+>", "", z.read(n).decode("utf8", "ignore")).strip()
            if contenido:
                con_texto.append(n.split("/")[-1])
        if con_texto:
            res.append((FALLA, f"Hay texto en encabezado/pie de página ({', '.join(con_texto)}). "
                               "Varios ATS lo ignoran: si ahí están tus datos de contacto, se pierden."))
        if not ("<w:txbxContent" in xml) and not con_texto:
            res.append((OK, "Sin cuadros de texto ni contenido en encabezado/pie."))
    return "\n".join(lineas)


# ----------------------------------------------------------------------------
# Revisiones sobre el texto extraído
# ----------------------------------------------------------------------------
SECCIONES = {
    "Experiencia": ["experiencia profesional", "experiencia laboral", "experiencia"],
    "Educación": ["educacion", "formacion academica", "formacion", "antecedentes academicos"],
    "Habilidades": ["habilidades", "competencias", "conocimientos", "herramientas"],
}


def revisar_texto(texto, res):
    palabras = len(texto.split())
    if palabras < 80:
        res.append((FALLA, f"Solo se extrajeron {palabras} palabras. Probablemente el CV es una imagen o "
                           "escaneo: el ATS lo verá casi vacío."))
        return
    res.append((OK, f"Texto extraíble ({palabras} palabras)."))

    n = normalizar(texto)
    inicio = n[: max(400, len(n) // 6)]
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", n)
    if not email:
        res.append((FALLA, "No se encontró un correo electrónico."))
    elif email.group(0) not in inicio:
        res.append((AVISO, "El correo no está al inicio del documento."))
    else:
        res.append((OK, f"Correo detectado al inicio ({email.group(0)})."))
    tel = re.search(r"(\+?\d[\d\s().-]{7,}\d)", n)
    if not tel:
        res.append((AVISO, "No se encontró un teléfono."))
    elif tel.group(0) not in inicio:
        res.append((AVISO, "El teléfono no está al inicio del documento."))
    else:
        res.append((OK, "Teléfono detectado al inicio."))

    faltan = []
    for nombre, variantes in SECCIONES.items():
        if not any(re.search(rf"(^|\n)\s*{re.escape(v)}\b", n, re.M) or v in n for v in variantes):
            faltan.append(nombre)
    if faltan:
        res.append((AVISO, f"No se reconocieron estas secciones estándar: {', '.join(faltan)}."))
    else:
        res.append((OK, "Secciones estándar reconocibles (Experiencia, Educación, Habilidades)."))

    raros = set(re.findall(r"[ﬀ-ﬆ�-]", texto))
    if raros:
        res.append((AVISO, f"Caracteres que pueden leerse mal: {' '.join(sorted(raros))} "
                           "(ligaduras o símbolos de fuentes especiales)."))
    else:
        res.append((OK, "Sin ligaduras ni símbolos ilegibles."))


ORDEN_POR_DEFECTO = ["perfil", "experiencia", "otra_experiencia", "proyectos",
                     "educacion", "certificaciones", "habilidades", "idiomas", "extra"]


def entradas_esperadas(datos):
    """Una entrada por trabajo, proyecto o estudio, en el orden en que build_cv.py
    las imprime. Cada una tiene un ancla (el cargo/título), campos de encabezado y
    logros."""
    entradas = []
    for clave in datos.get("orden") or ORDEN_POR_DEFECTO:
        if clave not in ("experiencia", "proyectos", "educacion"):
            continue
        for it in datos.get(clave) or []:
            if clave == "experiencia":
                ancla, campos = it.get("cargo"), [it.get("empresa"), it.get("ubicacion")]
            elif clave == "educacion":
                ancla, campos = it.get("titulo"), [it.get("institucion"), it.get("ubicacion")]
            else:
                ancla, campos = it.get("nombre"), []
            campos += [it.get("inicio"), it.get("fin")]
            logros = [re.sub(r"\*\*", "", l)[:45] for l in it.get("logros") or []]
            if ancla:
                entradas.append({"ancla": ancla, "campos": [c for c in campos if c], "logros": logros})
    return entradas


def evaluar_orden(texto, entradas):
    """Lo que necesita un ATS para atribuir bien la información:
    1. los cargos/títulos aparecen en orden;
    2. empresa, ciudad y fechas de cada entrada quedan dentro de su bloque
       (entre el cargo anterior y el siguiente), sin filtrarse a otro trabajo;
    3. los logros de cada entrada aparecen dentro de su bloque y en orden.
    El orden exacto dentro del encabezado (p. ej. si la ciudad sale antes o
    después del cargo) no importa: cada lector lo reconstruye distinto."""
    n = normalizar(texto)
    problemas = []
    anclas, pos = [], 0
    for e in entradas:
        a = normalizar(e["ancla"])
        i = n.find(a, pos)
        if i == -1:
            problemas.append(("falta" if n.find(a) == -1 else "desorden", e["ancla"]))
            anclas.append(None)
            continue
        anclas.append(i)
        pos = i + len(a)
    for k, e in enumerate(entradas):
        if anclas[k] is None:
            continue
        previas = [x for x in anclas[:k] if x is not None]
        siguientes = [x for x in anclas[k + 1:] if x is not None]
        lo = previas[-1] + 1 if previas else 0
        hi = siguientes[0] if siguientes else len(n)
        for c in e["campos"]:
            if n.find(normalizar(c), lo, hi + len(normalizar(c))) == -1:
                problemas.append(("fuera de su bloque" if n.find(normalizar(c)) != -1 else "falta", c))
        p = lo
        for l in e["logros"]:
            ln = normalizar(l)
            i = n.find(ln, p, hi)
            if i == -1:
                problemas.append(("fuera de su bloque" if n.find(ln) != -1 else "falta", l))
                continue
            p = i + 1
    return problemas


def revisar_orden(textos, datos, res):
    """Comprueba con cada lector que lo del JSON aparezca completo y bien atribuido."""
    entradas = entradas_esperadas(datos)
    if not entradas:
        return
    n_elem = sum(1 + len(e["campos"]) + len(e["logros"]) for e in entradas)
    fallidos = {}
    for lector, texto in textos.items():
        problemas = evaluar_orden(texto, entradas)
        if problemas:
            fallidos[lector] = problemas
    total = len(textos)
    if not fallidos:
        res.append((OK, f"Lectura correcta con {total} lector(es) distinto(s): los {n_elem} elementos "
                        "(cargos, empresas, fechas, logros) aparecen completos y cada uno en el bloque de "
                        "su trabajo o estudio."))
        return
    detalle = []
    for lector, probs in fallidos.items():
        tipo, ej = probs[0]
        detalle.append(f"{lector}: {len(probs)} problema(s), p. ej. «{ej[:45]}» ({tipo})")
    marca = FALLA if len(fallidos) * 2 > total or total == 1 else AVISO
    res.append((marca, f"Lectura con problemas en {len(fallidos)} de {total} lectores: " + "; ".join(detalle)
                       + ". El ATS podría asignar fechas o logros al trabajo equivocado."))


def revisar_claves(texto, claves, res):
    n = normalizar(texto)
    presentes, ausentes = [], []
    for k in claves:
        kn = normalizar(k)
        if not kn:
            continue
        # acepta plural simple (KPI → KPIs, gestor → gestores)
        patron = rf"(?<![\w]){re.escape(kn)}(?:s|es)?(?![\w])"
        (presentes if re.search(patron, n) else ausentes).append(k)
    total = len(presentes) + len(ausentes)
    if not total:
        return None
    res.append((OK if not ausentes else AVISO,
                f"Palabras clave de la oferta: {len(presentes)} de {total} presentes."))
    return presentes, ausentes


def main():
    ap = argparse.ArgumentParser(description="Verificación ATS de un CV (PDF o DOCX).")
    ap.add_argument("archivo")
    ap.add_argument("--json", help="JSON usado para generar el CV (activa la verificación de orden)")
    ap.add_argument("--keywords", help="Palabras clave separadas por coma")
    ap.add_argument("--keywords-file", help="Archivo con una palabra clave por línea")
    args = ap.parse_args()

    res = []
    if args.archivo.lower().endswith(".pdf"):
        textos = textos_pdf(args.archivo)
        revisar_pdf_tecnico(args.archivo, res)
    elif args.archivo.lower().endswith(".docx"):
        textos = {"docx": texto_docx(args.archivo, res)}
    else:
        sys.exit("Formato no soportado: usa .pdf o .docx")
    # para las revisiones generales usa el texto más completo
    texto = max(textos.values(), key=lambda t: len(t.split()))

    revisar_texto(texto, res)
    if args.json:
        with open(args.json, encoding="utf-8") as f:
            revisar_orden(textos, json.load(f), res)

    claves = []
    if args.keywords:
        claves += [k.strip() for k in args.keywords.split(",")]
    if args.keywords_file:
        with open(args.keywords_file, encoding="utf-8") as f:
            claves += [l.strip() for l in f if l.strip() and not l.startswith("#")]
    resultado_claves = revisar_claves(texto, claves, res) if claves else None

    print(f"Verificación ATS: {args.archivo}\n")
    for marca, msg in res:
        print(f"  {marca} {msg}")
    if resultado_claves:
        presentes, ausentes = resultado_claves
        print("\n  Presentes: " + (", ".join(presentes) or "—"))
        print("  Faltan:    " + (", ".join(ausentes) or "—"))
    fallas = sum(1 for m, _ in res if m == FALLA)
    avisos = sum(1 for m, _ in res if m == AVISO)
    print(f"\nResumen: {fallas} falla(s), {avisos} aviso(s).")
    sys.exit(1 if fallas else 0)


if __name__ == "__main__":
    main()
